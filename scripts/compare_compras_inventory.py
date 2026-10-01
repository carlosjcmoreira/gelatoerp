#!/usr/bin/env python3
"""Create a read-only reconciliation workbook for the Compras inventory.

The catalogue snapshot is exported separately using SELECT-only queries and is
passed as JSON. This script never connects to a database or changes source
workbooks or catalogue data.
"""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


STATUS_EXACT = "Correspondência exata candidata"
STATUS_REVIEW = "Revisão necessária"
STATUS_NONE = "Sem correspondência exata"
OUTPUT_HEADERS = [
    "Estado da comparação",
    "Fornecedor canónico candidato (ID)",
    "Fornecedor canónico candidato (nome)",
    "Evidência do fornecedor",
    "Artigo(s) candidato(s) (ID)",
    "Categoria atual candidata",
    "Ativo atual candidato",
    "Motivo / regra aplicada",
]
HISTORICAL_HEADERS = {
    "Valor Kg/L/UnValor total/sem iva",
    "QT Total 01.01",
    "Quant Balan.",
    "Valor Total 01.01",
    "Mat. 15 Jan",
    "Bolhão",
    "Garagem",
    "QT Total",
    "Valor Total",
}


def normalize_exact(value: Any) -> str:
    """Normalize case and whitespace only; preserve punctuation and accents."""
    if value is None:
        return ""
    return " ".join(str(value).split()).casefold()


def normalize_unit(value: Any) -> str:
    """Allow case/spacing variants only; do not convert units."""
    return normalize_exact(value)


def _supplier_lookup(
    suppliers: list[dict[str, Any]],
    aliases: list[dict[str, Any]],
) -> dict[str, dict[int, set[str]]]:
    lookup: dict[str, dict[int, set[str]]] = defaultdict(
        lambda: defaultdict(set)
    )
    for supplier in suppliers:
        supplier_id = supplier["id"]
        name = supplier.get("name")
        common_name = supplier.get("common_name")
        if normalize_exact(name):
            lookup[normalize_exact(name)][supplier_id].add("Nome canónico")
        if normalize_exact(common_name):
            lookup[normalize_exact(common_name)][supplier_id].add("Nome comum")

    for alias in aliases:
        # The export contains only aliases created during a confirmed merge.
        if alias.get("source") != "merge":
            continue
        key = normalize_exact(alias.get("alias_name"))
        if key:
            lookup[key][alias["supplier_id"]].add("Alias de fusão confirmada")
    return lookup


def _article_description(article: dict[str, Any]) -> str:
    supplier_name = (
        article.get("fornecedor_oficial_nome") or "sem fornecedor confirmado"
    )
    unit = article.get("unidade") or "unidade não registada"
    return (
        f"#{article['id']} — {supplier_name} — "
        f"{article.get('produto') or ''} — {unit}"
    )


def compare_inventory_row(
    row: dict[str, Any],
    articles: list[dict[str, Any]],
    suppliers: list[dict[str, Any]],
    aliases: list[dict[str, Any]],
) -> dict[str, Any]:
    """Compare one workbook row to the snapshot; never makes a merge decision."""
    supplier_label = row.get("Fornecedor")
    product = row.get("Produto")
    source_unit = row.get("Un")
    supplier_matches = _supplier_lookup(suppliers, aliases).get(
        normalize_exact(supplier_label), {}
    )
    supplier_ids = sorted(supplier_matches)
    supplier_by_id = {supplier["id"]: supplier for supplier in suppliers}
    matching_suppliers = [supplier_by_id[sid] for sid in supplier_ids]
    supplier_names = sorted(
        {
            supplier.get("name") or ""
            for supplier in matching_suppliers
        }
    )
    supplier_evidence = sorted(
        {
            evidence
            for evidence_set in supplier_matches.values()
            for evidence in evidence_set
        }
    )

    product_key = normalize_exact(product)
    product_candidates = [
        article
        for article in articles
        if normalize_exact(article.get("produto")) == product_key
    ]
    same_supplier_product = [
        article
        for article in product_candidates
        if len(supplier_ids) == 1
        and article.get("fornecedor_oficial_id") == supplier_ids[0]
    ]
    exact_candidates = [
        article
        for article in same_supplier_product
        if normalize_unit(article.get("unidade"))
        and normalize_unit(article.get("unidade")) == normalize_unit(source_unit)
    ]

    reasons: list[str] = []
    if not normalize_exact(supplier_label):
        reasons.append("O fornecedor está vazio na linha do inventário.")
    elif not supplier_ids:
        reasons.append(
            "O texto do fornecedor não coincide exatamente com nome, nome "
            "comum ou alias de fusão confirmada do cadastro."
        )
    elif len(supplier_ids) > 1:
        reasons.append(
            "O nome do fornecedor corresponde a mais de um registo canónico."
        )

    if not normalize_unit(source_unit):
        reasons.append("A unidade não está indicada no inventário.")

    if supplier_ids and len(supplier_ids) == 1 and product_key:
        if len(exact_candidates) == 1 and normalize_unit(source_unit):
            reasons.append(
                "Produto normalizado, fornecedor canónico e unidade coincidem "
                "com um único artigo."
            )
            status = STATUS_EXACT
            candidates = exact_candidates
        elif len(exact_candidates) > 1:
            reasons.append(
                "Existem vários artigos com produto, fornecedor e unidade "
                "iguais; é necessária escolha humana."
            )
            status = STATUS_REVIEW
            candidates = exact_candidates
        elif same_supplier_product:
            reasons.append(
                "O produto e fornecedor coincidem, mas a unidade está ausente "
                "ou é diferente no catálogo."
            )
            status = STATUS_REVIEW
            candidates = same_supplier_product
        elif product_candidates:
            reasons.append(
                "O produto existe com o mesmo nome, mas o fornecedor oficial "
                "e/ou a unidade não coincidem."
            )
            status = STATUS_REVIEW
            candidates = product_candidates
        else:
            reasons.append(
                "Não existe produto com nome normalizado exatamente igual "
                "no catálogo atual."
            )
            status = STATUS_NONE
            candidates = []
    else:
        if product_candidates:
            reasons.append(
                "O produto existe no catálogo, mas os dados do fornecedor "
                "não permitem confirmar a correspondência."
            )
            candidates = product_candidates
        else:
            reasons.append(
                "Não foi possível confirmar uma correspondência porque os "
                "dados de fornecedor ou unidade estão incompletos."
            )
            candidates = []
        status = STATUS_REVIEW

    candidate_ids = sorted({article["id"] for article in candidates})
    categories = sorted(
        {
            article.get("categoria_artigo") or ""
            for article in candidates
            if article.get("categoria_artigo")
        }
    )
    active_values = sorted(
        {article.get("ativo") for article in candidates if article.get("ativo") is not None}
    )
    return {
        "status": status,
        "supplier_ids": supplier_ids,
        "supplier_names": supplier_names,
        "supplier_evidence": supplier_evidence,
        "article_ids": candidate_ids,
        "candidate_descriptions": [
            _article_description(article)
            for article in sorted(candidates, key=lambda item: item["id"])
        ],
        "categories": categories,
        "active_values": active_values,
        "reason": " ".join(reasons),
    }


def read_inventory(path: Path) -> tuple[str, list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers_row = next(values, None)
        if not headers_row:
            raise ValueError("O inventário não tem uma linha de cabeçalhos.")
        headers = [
            str(header).strip() if header is not None else f"Coluna {index + 1}"
            for index, header in enumerate(headers_row)
        ]
        required = {"Fornecedor", "Produto", "Un"}
        missing = sorted(required - set(headers))
        if missing:
            raise ValueError(
                "Faltam colunas necessárias no inventário: "
                + ", ".join(missing)
            )

        rows: list[dict[str, Any]] = []
        for excel_row, values_row in enumerate(values, start=2):
            row = {
                headers[index]: values_row[index] if index < len(values_row) else None
                for index in range(len(headers))
            }
            if normalize_exact(row.get("Produto")):
                row["__source_row"] = excel_row
                rows.append(row)
        return sheet.title, headers, rows
    finally:
        workbook.close()


def _display_join(values: list[Any]) -> str:
    return "; ".join(str(value) for value in values if value not in (None, ""))


def _prepare_inventory_rows(
    headers: list[str],
    source_rows: list[dict[str, Any]],
    snapshot: dict[str, Any],
) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for source in source_rows:
        comparison = compare_inventory_row(
            source,
            snapshot["articles"],
            snapshot["suppliers"],
            snapshot.get("confirmed_merge_aliases", []),
        )
        raw_values = []
        for header in headers:
            output_header = (
                f"{header} [valor histórico do ficheiro]"
                if header in HISTORICAL_HEADERS
                else header
            )
            raw_values.append(source.get(header))
        output.append(
            {
                "source_row": source["__source_row"],
                "raw_values": raw_values,
                "comparison": comparison,
            }
        )
    return output


def _write_table(ws, headers: list[str], rows: list[list[Any]], table_name: str) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if rows:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="23405D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 38
    for column in ws.columns:
        letter = column[0].column_letter
        width = max(
            (len(str(cell.value)) for cell in column[: min(len(column), 200)] if cell.value is not None),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(width + 2, 12), 42)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def create_workbook(
    inventory_path: Path,
    snapshot_path: Path,
    output_path: Path,
    snapshot_date: str | None = None,
) -> dict[str, Any]:
    with snapshot_path.open("r", encoding="utf-8") as handle:
        snapshot = json.load(handle)
    if snapshot.get("environment") != "production":
        raise ValueError(
            "O snapshot tem de vir da base publicada; não substitua por desenvolvimento."
        )
    if not {"articles", "suppliers"} <= snapshot.keys():
        raise ValueError("O snapshot não inclui artigos e fornecedores.")

    sheet_name, source_headers, source_rows = read_inventory(inventory_path)
    prepared = _prepare_inventory_rows(source_headers, source_rows, snapshot)
    counts = Counter(item["comparison"]["status"] for item in prepared)

    workbook = Workbook()
    compared = workbook.active
    compared.title = "Inventário comparado"
    historical_headers = [
        f"{header} [valor histórico do ficheiro]"
        if header in HISTORICAL_HEADERS
        else header
        for header in source_headers
    ]
    inventory_headers = [
        "Linha no inventário",
        *historical_headers,
        *OUTPUT_HEADERS,
        "Descrição dos artigos candidatos",
    ]
    inventory_data: list[list[Any]] = []
    for item in prepared:
        comparison = item["comparison"]
        inventory_data.append(
            [
                item["source_row"],
                *item["raw_values"],
                comparison["status"],
                _display_join(comparison["supplier_ids"]),
                _display_join(comparison["supplier_names"]),
                _display_join(comparison["supplier_evidence"]),
                _display_join(comparison["article_ids"]),
                _display_join(comparison["categories"]),
                _display_join(
                    ["Ativo" if value else "Inativo" for value in comparison["active_values"]]
                ),
                comparison["reason"],
                _display_join(comparison["candidate_descriptions"]),
            ]
        )
    _write_table(compared, inventory_headers, inventory_data, "InventarioComparado")
    status_column = len(inventory_headers) - len(OUTPUT_HEADERS)
    status_letter = compared.cell(row=1, column=status_column).column_letter
    last_row = max(2, compared.max_row)
    for status, color in (
        (STATUS_EXACT, "E2F0D9"),
        (STATUS_REVIEW, "FFF2CC"),
        (STATUS_NONE, "DDEBF7"),
    ):
        compared.conditional_formatting.add(
            f"{status_letter}2:{status_letter}{last_row}",
            FormulaRule(
                formula=[f'${status_letter}2="{status}"'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )

    catalogue = workbook.create_sheet("Catálogo atual")
    catalogue_headers = [
        "ID",
        "Fornecedor (etiqueta atual)",
        "Produto",
        "Marca",
        "Unidade",
        "Fornecedor oficial (ID)",
        "Fornecedor oficial (nome)",
        "Categoria atual",
        "Ativo atual",
        "Conjunto de origem",
        "Versão de origem",
        "Linha da origem",
        "Etiqueta histórica original",
        "Origem operacional atual",
        "Tipo de origem operacional",
    ]
    catalogue_data = []
    for article in sorted(snapshot["articles"], key=lambda item: item["id"]):
        catalogue_data.append(
            [
                article.get("id"),
                article.get("fornecedor"),
                article.get("produto"),
                article.get("marca"),
                article.get("unidade"),
                article.get("fornecedor_oficial_id"),
                article.get("fornecedor_oficial_nome"),
                article.get("categoria_artigo"),
                "Ativo" if article.get("ativo") else "Inativo",
                article.get("source_dataset"),
                article.get("source_version"),
                article.get("source_row"),
                article.get("origem_original"),
                article.get("origem_nome"),
                article.get("origem_tipo"),
            ]
        )
    _write_table(catalogue, catalogue_headers, catalogue_data, "CatalogoAtual")

    summary = workbook.create_sheet("Resumo e decisões")
    summary_rows: list[list[Any]] = [
        ["Reconciliação do inventário com o Catálogo de Compras", ""],
        ["Inventário de origem", inventory_path.name],
        ["Folha do Excel de origem", sheet_name],
        ["Linhas com produto", len(source_rows)],
        ["Snapshot do catálogo", snapshot.get("as_of", snapshot_date or "Data não indicada")],
        ["Artigos no catálogo", len(snapshot["articles"])],
        ["Fornecedores no cadastro consultados", len(snapshot["suppliers"])],
        ["Aliases considerados", len(snapshot.get("confirmed_merge_aliases", []))],
        ["", ""],
        ["Estado da comparação", "Número de linhas"],
        [STATUS_EXACT, counts.get(STATUS_EXACT, 0)],
        [STATUS_REVIEW, counts.get(STATUS_REVIEW, 0)],
        [STATUS_NONE, counts.get(STATUS_NONE, 0)],
        ["", ""],
        ["Como ler os resultados", ""],
        [
            "Correspondência exata candidata",
            "Um único produto com nome, fornecedor canónico e unidade exatamente compatíveis após normalização de maiúsculas/minúsculas e espaços. É uma sugestão para validar, não uma associação gravada.",
        ],
        [
            "Revisão necessária",
            "Falta informação, existe ambiguidade, a unidade não coincide ou o produto existe ligado a outro fornecedor.",
        ],
        [
            "Sem correspondência exata",
            "Nenhum produto com nome normalizado exatamente igual foi encontrado; isto não confirma que seja um artigo novo.",
        ],
        [
            "Normalização",
            "Só ignora maiúsculas/minúsculas e espaços repetidos. Não ignora acentos, pontuação, tamanhos, marcas ou palavras; não converte unidades.",
        ],
        [
            "Fornecedor",
            "Usa nomes canónicos, nomes comuns e aliases registados após fusões confirmadas. Nunca cria nem confirma fornecedores.",
        ],
        [
            "Atenção às quantidades e valores",
            f"O ficheiro de origem contém uma folha chamada «{sheet_name}» e colunas com datas/locais distintos. Os valores são reproduzidos apenas como dados históricos do ficheiro, não são stock nem preços atuais.",
        ],
        ["", ""],
        ["Decisões de master data para validar", ""],
        [
            "Ficha proposta",
            "Manter identidade do artigo, categoria de master data, estado ativo e elegibilidade para encomenda pelas lojas como conceitos separados. A elegibilidade ainda não existe como campo independente no catálogo atual.",
        ],
        [
            "Transição recomendada",
            "Manter a disponibilidade atual dos artigos existentes; artigos novos só ficam elegíveis para encomenda depois de validados.",
        ],
        [
            "Fornecedores alternativos",
            "Decidir se o mesmo artigo pode ter mais de um fornecedor oficial. A comparação não funde artigos nem substitui associações atuais.",
        ],
        [
            "Contagem de artigos não encomendáveis",
            "Decidir se as lojas podem continuar a contar artigos ativos que não podem encomendar.",
        ],
        [
            "Categorias",
            "Preservar categorias e decisões existentes. Não classificar automaticamente artigos novos com base no fornecedor, na origem, no preço ou no nome.",
        ],
        ["", ""],
        ["Fontes e limites", ""],
        [
            "Inventário",
            "Ficheiro fornecido pela equipa; os valores originais são mantidos na folha «Inventário comparado».",
        ],
        [
            "Catálogo e fornecedores",
            f"Snapshot de leitura da base publicada em {snapshot.get('as_of', snapshot_date or 'data não indicada')}; campos empresariais mínimos, sem NIF, IBAN, contactos ou notas.",
        ],
        [
            "Não alterado",
            "Este ficheiro não importa artigos nem altera fornecedores, artigos, categorias, encomendas, contagens, stocks ou documentos históricos.",
        ],
    ]
    for row in summary_rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 37
    summary.column_dimensions["B"].width = 105
    summary.freeze_panes = "A2"
    summary["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    summary["B1"].fill = PatternFill("solid", fgColor="23405D")
    summary["A1"].fill = PatternFill("solid", fgColor="23405D")
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in (10, 15, 23, 30):
        for cell in summary[row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="23405D")
    summary.row_dimensions[1].height = 26
    for row_num in range(2, summary.max_row + 1):
        summary.row_dimensions[row_num].height = 34

    workbook.properties.title = "Reconciliação do inventário com o Catálogo de Compras"
    workbook.properties.subject = "Comparação somente de leitura para revisão humana"
    workbook.properties.creator = "Replit Agent"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return {
        "inventory_rows": len(source_rows),
        "catalogue_articles": len(snapshot["articles"]),
        "status_counts": dict(counts),
        "output": str(output_path),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--inventory", type=Path, required=True)
    parser.add_argument("--snapshot", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = create_workbook(args.inventory, args.snapshot, args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()