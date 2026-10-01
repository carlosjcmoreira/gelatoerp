import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.compare_compras_inventory import (
    STATUS_EXACT,
    STATUS_NONE,
    STATUS_REVIEW,
    compare_inventory_row,
    create_workbook,
)


class CompareComprasInventoryTests(unittest.TestCase):
    def setUp(self):
        self.suppliers = [{"id": 10, "name": "Fornecedor Exemplo", "common_name": None}]
        self.aliases = []

    def article(self, article_id=1, product="Produto A", unit="Kg", supplier_id=10):
        return {
            "id": article_id,
            "produto": product,
            "unidade": unit,
            "fornecedor_oficial_id": supplier_id,
            "fornecedor_oficial_nome": "Fornecedor Exemplo",
            "fornecedor": "Fornecedor Exemplo",
            "categoria_artigo": "Matéria-prima",
            "ativo": True,
        }

    def row(self, supplier="fornecedor exemplo", product="Produto A", unit="kg"):
        return {"Fornecedor": supplier, "Produto": product, "Un": unit}

    def test_unique_exact_match_ignores_only_case_and_spacing(self):
        result = compare_inventory_row(
            self.row(supplier="  FORNECEDOR   EXEMPLO ", product=" Produto A ", unit="KG"),
            [self.article()],
            self.suppliers,
            self.aliases,
        )
        self.assertEqual(result["status"], STATUS_EXACT)
        self.assertEqual(result["article_ids"], [1])

    def test_multiple_exact_candidates_require_review(self):
        result = compare_inventory_row(
            self.row(),
            [self.article(1), self.article(2)],
            self.suppliers,
            self.aliases,
        )
        self.assertEqual(result["status"], STATUS_REVIEW)
        self.assertEqual(result["article_ids"], [1, 2])

    def test_missing_supplier_requires_review(self):
        result = compare_inventory_row(
            self.row(supplier=None),
            [self.article()],
            self.suppliers,
            self.aliases,
        )
        self.assertEqual(result["status"], STATUS_REVIEW)
        self.assertEqual(result["article_ids"], [1])

    def test_incompatible_unit_requires_review(self):
        result = compare_inventory_row(
            self.row(unit="Un"),
            [self.article(unit="Kg")],
            self.suppliers,
            self.aliases,
        )
        self.assertEqual(result["status"], STATUS_REVIEW)
        self.assertEqual(result["article_ids"], [1])

    def test_composite_supplier_label_is_explicitly_flagged(self):
        result = compare_inventory_row(
            self.row(supplier="Fornecedor Exemplo / outro"),
            [self.article()],
            self.suppliers,
            self.aliases,
        )
        self.assertEqual(result["status"], STATUS_REVIEW)
        self.assertIn("combina mais de um fornecedor", result["reason"])

    def test_workbook_preserves_all_product_rows_and_pending_cases(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            inventory_path = temp / "inventory.xlsx"
            snapshot_path = temp / "snapshot.json"
            output_path = temp / "comparison.xlsx"

            source = Workbook()
            sheet = source.active
            sheet.title = "Inventory"
            sheet.append(["Fornecedor", "Produto", "Un", "Quantidade"])
            sheet.append(["Fornecedor Exemplo", "Produto A", "Kg", 5])
            sheet.append([None, "Produto A", "Kg", 6])
            sheet.append(["Fornecedor Exemplo", "Produto B", "L", 7])
            sheet.append(["Fornecedor Exemplo", "Novo Produto", "Un", 8])
            sheet.append(["Fornecedor Exemplo", None, "Un", 999])
            source.save(inventory_path)
            source.close()

            snapshot = {
                "environment": "production",
                "as_of": "2026-10-01",
                "articles": [
                    self.article(1, product="Produto A", unit="Kg"),
                    self.article(2, product="Produto B", unit="Kg"),
                ],
                "suppliers": self.suppliers,
                "confirmed_merge_aliases": [],
            }
            snapshot_path.write_text(
                json.dumps(snapshot, ensure_ascii=False), encoding="utf-8"
            )

            result = create_workbook(inventory_path, snapshot_path, output_path)
            self.assertEqual(result["inventory_rows"], 4)
            self.assertEqual(result["status_counts"][STATUS_EXACT], 1)
            self.assertEqual(result["status_counts"][STATUS_REVIEW], 2)
            self.assertEqual(result["status_counts"][STATUS_NONE], 1)

            output = load_workbook(output_path, read_only=True, data_only=True)
            self.assertEqual(
                output.sheetnames,
                ["Inventário comparado", "Catálogo atual", "Resumo e decisões"],
            )
            compared = output["Inventário comparado"]
            data_rows = list(compared.iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(data_rows), 4)
            self.assertEqual([row[0] for row in data_rows], [2, 3, 4, 5])
            summary_values = [
                cell.value
                for row in output["Resumo e decisões"].iter_rows()
                for cell in row
            ]
            self.assertIn("Linhas que precisam de validação", summary_values)
            output.close()


if __name__ == "__main__":
    unittest.main()