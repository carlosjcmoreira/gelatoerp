"""Cost-center filtering and display for the Contabilidade invoice list."""

import unittest
from datetime import date
from io import BytesIO
from unittest.mock import patch

from flask import Blueprint, Flask
from jinja2 import ChoiceLoader, DictLoader, FileSystemLoader


class _RecordingCursor:
    def __init__(self, rows=(), one=None):
        self.rows = rows
        self.one = one
        self.calls = []

    def execute(self, sql, params=None):
        self.calls.append((sql, params))

    def fetchall(self):
        return self.rows

    def fetchone(self):
        return self.one


class _RecordingConnection:
    def __init__(self, cursor):
        self._cursor = cursor

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False

    def cursor(self):
        return self._cursor


def _make_app():
    from flask_app.routes.contabilidade import contabilidade_bp

    app = Flask(__name__)
    app.secret_key = 'test-secret-key'
    app.config['TESTING'] = True

    auth_bp = Blueprint('auth', __name__)

    @auth_bp.route('/login')
    def login():
        return 'login'

    home_bp = Blueprint('home', __name__)

    @home_bp.route('/')
    def index():
        return 'home'

    app.register_blueprint(auth_bp)
    app.register_blueprint(home_bp)
    app.register_blueprint(contabilidade_bp, url_prefix='/contabilidade')
    return app


class TestCostCenterWhereClause(unittest.TestCase):
    def test_one_center_matches_legacy_or_split_assignments_without_joining_rows(self):
        from db.contabilidade import _build_cont_where

        where, params = _build_cont_where(
            supplier_name='Fornecedor',
            accounting_status='contabilizado',
            centro_custo_id=17,
            search='FT-42',
        )

        self.assertIn('i.centro_custo_id = %s', where)
        self.assertIn('EXISTS (', where)
        self.assertIn('FROM invoice_centros_custo icc', where)
        self.assertEqual(params, [
            '%Fornecedor%', 'contabilizado', 17, 17,
            '%FT-42%', '%FT-42%', '%FT-42%',
        ])

    def test_unassigned_means_no_legacy_or_split_center(self):
        from db.contabilidade import _build_cont_where

        where, params = _build_cont_where(sem_cc=True)

        self.assertIn('i.centro_custo_id IS NULL', where)
        self.assertIn('NOT EXISTS', where)
        self.assertNotIn('i.centro_custo_id = %s', where)
        self.assertEqual(params, [])

    def test_invoice_query_displays_split_centers_and_uses_unique_invoice_rows(self):
        from db import contabilidade

        row = (
            11, 4, 'Fornecedor', 'PT501234567', 'FT-11', 'fatura',
            120.0, 23.0, date(2026, 5, 1), 'pending_review', None,
            'por_contabilizar', None, None, None, None, True,
            'Fornecedor Legal', 'Fornecedor', 'Matosinhos, Bolhão',
        )
        cursor = _RecordingCursor(rows=[row])
        connection = _RecordingConnection(cursor)

        with patch.object(contabilidade, 'db_connection', return_value=connection):
            invoices = contabilidade.get_cont_invoices(centro_custo_id=17)

        self.assertEqual(len(invoices), 1)
        self.assertEqual(invoices[0]['centro_custo_name'], 'Matosinhos, Bolhão')
        sql, params = cursor.calls[0]
        self.assertIn('STRING_AGG(cc2.name', sql)
        self.assertIn('LEFT JOIN cost_centers cc ON cc.id = i.centro_custo_id', sql)
        self.assertEqual(sql.count('FROM invoice_centros_custo icc2'), 1)
        self.assertNotIn('JOIN invoice_centros_custo icc2 ON', sql)
        self.assertEqual(params, [17, 17, 50, 0])

    def test_count_uses_the_same_no_center_filter_without_join_multiplication(self):
        from db import contabilidade

        cursor = _RecordingCursor(one=(3,))
        connection = _RecordingConnection(cursor)

        with patch.object(contabilidade, 'db_connection', return_value=connection):
            count = contabilidade.count_cont_invoices(sem_cc=True)

        self.assertEqual(count, 3)
        sql, params = cursor.calls[0]
        self.assertIn('COUNT(*) FROM invoices i', sql)
        self.assertIn('NOT EXISTS', sql)
        self.assertEqual(params, [])


class TestContabilidadeCostCenterRoutes(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = _make_app()
        cls.app.config['TESTING'] = True

    def setUp(self):
        self.client = self.app.test_client()
        with self.client.session_transaction() as session:
            session['user'] = {
                'id': 7,
                'username': 'contabilidade-tester',
                'acesso_contabilidade': True,
            }

    def _patch_list_route(self, count=0):
        return (
            patch('db.contabilidade.count_cont_invoices', return_value=count),
            patch('db.contabilidade.get_cont_invoices', return_value=[]),
            patch('db.contabilidade.get_cont_summary', return_value={
                'por_contabilizar_count': 0,
                'por_contabilizar_eur': 0,
                'contabilizado_mes': 0,
                'tickets_abertos': 0,
            }),
            patch('db.faturas.get_stores_list', return_value=[]),
            patch('db.faturas.get_distinct_supplier_names', return_value=[]),
            patch('db.centros_custo.get_cost_centers', return_value=[
                {'id': 17, 'code': 'MAT', 'name': 'Matosinhos', 'ativo': True},
                {'id': 18, 'code': 'BOL', 'name': 'Bolhão', 'ativo': False},
            ]),
            patch('flask_app.routes.contabilidade.render_template', return_value='ok'),
        )

    def test_list_passes_cost_center_and_other_filters_to_count_and_rows(self):
        patches = self._patch_list_route(count=120)
        with patches[0] as count, patches[1] as get_rows:
            with patches[2], patches[3], patches[4], patches[5], patches[6] as render:
                response = self.client.get(
                    '/contabilidade/?centro_custo_id=17&q=farinha'
                    '&accounting_status=por_contabilizar&date_from=2026-01-01&page=2'
                )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(count.call_args.kwargs['centro_custo_id'], 17)
        self.assertFalse(count.call_args.kwargs['sem_cc'])
        self.assertEqual(count.call_args.kwargs['search'], 'farinha')
        self.assertEqual(count.call_args.kwargs['accounting_status'], 'por_contabilizar')
        self.assertEqual(count.call_args.kwargs['date_from'], date(2026, 1, 1))
        self.assertEqual(get_rows.call_args.kwargs['centro_custo_id'], 17)
        self.assertEqual(get_rows.call_args.kwargs['offset'], 50)
        self.assertEqual(render.call_args.kwargs['centro_custo_filter'], '17')
        self.assertTrue(render.call_args.kwargs['has_filters'])
        self.assertEqual(
            [cc['id'] for cc in render.call_args.kwargs['cost_centers']],
            [17, 18],
        )

    def test_list_maps_none_option_to_sem_cc_filter(self):
        patches = self._patch_list_route()
        with patches[0] as count:
            with patches[1], patches[2], patches[3], patches[4], patches[5], patches[6] as render:
                response = self.client.get('/contabilidade/?centro_custo_id=__none__')

        self.assertEqual(response.status_code, 200)
        self.assertIsNone(count.call_args.kwargs['centro_custo_id'])
        self.assertTrue(count.call_args.kwargs['sem_cc'])
        self.assertEqual(render.call_args.kwargs['centro_custo_filter'], '__none__')

    def test_rendered_list_displays_centers_and_preserves_filters_in_links(self):
        invoice = {
            'id': 11,
            'issue_date': date(2026, 5, 1),
            'supplier_legal_name': 'Fornecedor Legal',
            'supplier_display_name': 'Fornecedor',
            'supplier_name': 'Fornecedor',
            'supplier_nif': 'PT501234567',
            'centro_custo_name': 'Matosinhos, Bolhão',
            'invoice_number': 'FT-11',
            'document_type': 'fatura',
            'amount_eur': 120.0,
            'vat_amount_eur': 23.0,
            'status': 'pending_review',
            'accounting_status': 'contabilizado',
            'accounting_notes': None,
            'has_pdf': False,
        }
        self.app.jinja_loader = ChoiceLoader([
            DictLoader({'base.html': '{% block content %}{% endblock %}'}),
            FileSystemLoader('flask_app/templates'),
        ])
        with patch('db.contabilidade.count_cont_invoices', return_value=120), \
                patch('db.contabilidade.get_cont_invoices', return_value=[invoice]), \
                patch('db.contabilidade.get_cont_summary', return_value={
                    'por_contabilizar_count': 0,
                    'por_contabilizar_eur': 0,
                    'contabilizado_mes': 0,
                    'tickets_abertos': 0,
                }), \
                patch('db.faturas.get_stores_list', return_value=[]), \
                patch('db.faturas.get_distinct_supplier_names', return_value=['Fornecedor']), \
                patch('db.centros_custo.get_cost_centers', return_value=[
                    {'id': 17, 'code': 'MAT', 'name': 'Matosinhos', 'ativo': True},
                ]):
            response = self.client.get(
                '/contabilidade/?centro_custo_id=17&supplier_name=Fornecedor'
                '&document_type=fatura&accounting_status=contabilizado'
                '&date_from=2026-01-01&q=farinha&page=2'
            )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Matosinhos, Bolh', response.data)
        self.assertIn(b'name="centro_custo_id"', response.data)
        self.assertIn(b'value="17" selected', response.data)
        self.assertIn(b'centro_custo_id=17', response.data)
        self.assertIn(b'supplier_name=Fornecedor', response.data)
        self.assertIn(b'document_type=fatura', response.data)
        self.assertIn(b'accounting_status=contabilizado', response.data)
        self.assertIn(b'date_from=2026-01-01', response.data)
        self.assertIn(b'q=farinha', response.data)

    def test_excel_includes_cost_center_and_applies_same_filter(self):
        from openpyxl import load_workbook

        invoice = {
            'issue_date': date(2026, 5, 1),
            'supplier_display_name': 'Fornecedor',
            'supplier_name': 'Fornecedor',
            'supplier_nif': 'PT501234567',
            'centro_custo_name': 'Matosinhos, Bolhão',
            'invoice_number': 'FT-11',
            'document_type': 'fatura',
            'amount_eur': 120.0,
            'vat_amount_eur': 23.0,
            'status': 'pending_review',
            'accounting_status': 'por_contabilizar',
            'accounting_notes': None,
            'accounting_updated_at': None,
            'accounting_updated_by': None,
        }
        with patch('db.contabilidade.get_cont_invoices', return_value=[invoice]) as get_rows:
            response = self.client.get(
                '/contabilidade/export.xlsx?centro_custo_id=__none__'
                '&q=farinha&date_from=2026-01-01'
            )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.mimetype.endswith('spreadsheetml.sheet'))
        self.assertTrue(get_rows.call_args.kwargs['sem_cc'])
        self.assertIsNone(get_rows.call_args.kwargs['centro_custo_id'])
        self.assertEqual(get_rows.call_args.kwargs['search'], 'farinha')
        self.assertEqual(get_rows.call_args.kwargs['date_from'], date(2026, 1, 1))

        workbook = load_workbook(BytesIO(response.data), read_only=True)
        sheet = workbook['Contabilidade']
        self.assertEqual(sheet.cell(1, 4).value, 'Centro de Custo')
        self.assertEqual(sheet.cell(2, 4).value, 'Matosinhos, Bolhão')
        self.assertEqual(sheet.cell(2, 5).value, 'FT-11')
        workbook.close()


class TestContabilidadeCostCenterTemplate(unittest.TestCase):
    def test_column_filter_and_sort_pagination_links_preserve_center_selection(self):
        with open('flask_app/templates/contabilidade/index.html', encoding='utf-8') as template:
            source = template.read()

        self.assertIn('name="centro_custo_id"', source)
        self.assertIn('value="__none__"', source)
        self.assertIn('Sem centro de custo', source)
        self.assertIn('{{ inv.centro_custo_name or \'Sem centro de custo\' }}', source)
        self.assertIn('centro_custo_id=centro_custo_filter or None', source)
        self.assertIn('url_for(\'contabilidade.index\', page=target_page', source)
