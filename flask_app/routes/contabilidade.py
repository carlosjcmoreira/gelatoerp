"""Contabilidade blueprint.

Provides a unified view of all invoices for the accounting team, with:
- accounting_status management (single + bulk)
- Excel export respecting active filters
- Tickets system for tracking missing documents
"""
import logging
import math
import os
import re
import unicodedata
import zipfile
from datetime import date as _date, datetime
from io import BytesIO
from urllib.parse import quote

from flask import (
    Blueprint, render_template, request, redirect, url_for,
    flash, session, jsonify, send_file, abort,
)
from flask_app.auth import perm_required

logger = logging.getLogger(__name__)

contabilidade_bp = Blueprint('contabilidade', __name__)


def _username() -> str:
    return session.get('user', {}).get('username', 'sistema')


def _parse_date(raw: str):
    if not raw:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(raw.strip(), fmt).date()
        except ValueError:
            pass
    return None


def _safe_contabilidade_return_url(raw: str):
    """Accept only local Contabilidade URLs for post-ticket redirects."""
    from urllib.parse import urlparse

    value = (raw or '').strip()
    if not value or value.startswith('//') or '\\' in value:
        return None
    parsed = urlparse(value)
    if parsed.scheme or parsed.netloc:
        return None
    if parsed.path != '/contabilidade' and not parsed.path.startswith('/contabilidade/'):
        return None
    return value


def _parse_centro_custo_filter(raw: str):
    raw = (raw or '').strip()
    if raw == '__none__':
        return None, True, raw
    try:
        centro_custo_id = int(raw)
    except (TypeError, ValueError):
        return None, False, ''
    if centro_custo_id <= 0:
        return None, False, ''
    return centro_custo_id, False, str(centro_custo_id)


PAGE_SIZE = 50
MAX_CONTABILIDADE_ZIP_DOCUMENTS = PAGE_SIZE
MAX_CONTABILIDADE_ZIP_BYTES = 50 * 1024 * 1024


def _safe_center_filename_part(value: str) -> str:
    """Keep readable Unicode while removing path and header-sensitive characters."""
    normalized = unicodedata.normalize('NFC', str(value or '')).strip()
    safe = ''.join(
        char if char.isalnum() or char in '-_' else '_'
        for char in normalized
    )
    return re.sub(r'_+', '_', safe).strip('_-')


def _safe_invoice_download_name(filename, mimetype: str, cost_center_names=()) -> str:
    """Build a safe download name without changing the stored invoice filename."""
    original = str(filename or '')
    basename = original.replace('\\', '/').rsplit('/', 1)[-1]
    safe_chars = []
    for char in basename:
        if unicodedata.category(char).startswith('C') or char in '<>:"|?*;':
            safe_chars.append('_')
        else:
            safe_chars.append(char)
    basename = ''.join(safe_chars).strip().rstrip(' .')
    stem, extension = os.path.splitext(basename)
    stem = stem.strip().rstrip(' .') or 'fatura'

    if not re.fullmatch(r'\.[A-Za-z0-9]{1,10}', extension):
        extension = ''

    expected_extensions = {
        'application/pdf': ('.pdf', {'.pdf'}),
        'image/jpeg': ('.jpg', {'.jpg', '.jpeg'}),
        'image/png': ('.png', {'.png'}),
    }
    expected = expected_extensions.get(mimetype)
    if expected and extension.lower() not in expected[1]:
        extension = expected[0]
    elif not extension and not original:
        extension = expected[0] if expected else '.bin'

    suffixes = {}
    for name in cost_center_names or ():
        safe_name = _safe_center_filename_part(name)
        if safe_name:
            suffixes.setdefault(safe_name.casefold(), safe_name)
    suffix = '_'.join(sorted(suffixes.values(), key=lambda value: (value.casefold(), value)))
    if suffix:
        stem = f'{stem}_{suffix}'
    return f'{stem}{extension}'


def _safe_zip_folder_name(raw_name) -> str:
    """Normalize a user-supplied ZIP/folder name to a safe single path component."""
    if not isinstance(raw_name, str):
        raise ValueError('Indique um nome para o ZIP.')
    name = unicodedata.normalize('NFC', raw_name).strip()
    if name.casefold().endswith('.zip'):
        name = name[:-4].strip()
    if not name or len(name) > 80:
        raise ValueError('O nome do ZIP deve ter entre 1 e 80 caracteres.')

    safe = ''.join(
        char if char.isalnum() or char in ' _-' else '_'
        for char in name
    )
    safe = re.sub(r' +', ' ', safe)
    safe = re.sub(r'_+', '_', safe).strip(' _-.')
    if not safe or len(safe) > 80:
        raise ValueError('Indique um nome válido para o ZIP.')
    return safe


def _unique_zip_entry_filename(filename: str, invoice_id: int, used_names: set) -> str:
    """Disambiguate duplicate basenames without allowing path components."""
    stem, extension = os.path.splitext(filename)
    candidate = filename
    if candidate.casefold() in used_names:
        candidate = f'{stem}_{invoice_id}{extension}'
        suffix = 2
        while candidate.casefold() in used_names:
            candidate = f'{stem}_{invoice_id}_{suffix}{extension}'
            suffix += 1
    used_names.add(candidate.casefold())
    return candidate


def _detect_invoice_mime(data, filename):
    if isinstance(data, memoryview):
        data = bytes(data)
    if data.startswith(b'%PDF'):
        return 'application/pdf'
    if len(data) >= 2 and data[:2] == b'\xff\xd8':
        return 'image/jpeg'
    if len(data) >= 8 and data[:8] == b'\x89PNG\r\n\x1a\n':
        return 'image/png'
    ext = (filename or '').lower().rsplit('.', 1)[-1]
    return {
        'pdf': 'application/pdf',
        'jpg': 'image/jpeg',
        'jpeg': 'image/jpeg',
        'png': 'image/png',
    }.get(ext, 'application/octet-stream')


# ── Main listing ──────────────────────────────────────────────────────────────

@contabilidade_bp.route('/')
@perm_required('acesso_contabilidade')
def index():
    from db.contabilidade import (
        get_cont_invoices, count_cont_invoices, get_cont_summary,
        ACCOUNTING_STATUS_LABELS,
    )
    from db.faturas import get_stores_list, get_distinct_supplier_names, DOCUMENT_TYPE_LABELS
    from db.centros_custo import get_cost_centers

    today = _date.today()

    supplier_name = request.args.get('supplier_name', '').strip()
    document_type = request.args.get('document_type', '').strip()
    if document_type not in DOCUMENT_TYPE_LABELS:
        document_type = ''
    accounting_status = request.args.get('accounting_status', '').strip()
    if accounting_status not in ACCOUNTING_STATUS_LABELS and accounting_status != '':
        accounting_status = ''
    date_from = _parse_date(request.args.get('date_from', ''))
    date_to = _parse_date(request.args.get('date_to', ''))
    date_from_raw = request.args.get('date_from', '')
    date_to_raw = request.args.get('date_to', '')
    search = request.args.get('q', '').strip()
    centro_custo_id, sem_cc, centro_custo_filter = _parse_centro_custo_filter(
        request.args.get('centro_custo_id', '')
    )
    order_by = request.args.get('order_by', 'issue_date')
    order_dir = request.args.get('order_dir', 'desc')

    try:
        page = max(1, int(request.args.get('page', '1') or '1'))
    except ValueError:
        page = 1

    filter_kwargs = dict(
        supplier_name=supplier_name or None,
        document_type=document_type or None,
        accounting_status=accounting_status or None,
        date_from=date_from,
        date_to=date_to,
        search=search or None,
        centro_custo_id=centro_custo_id,
        sem_cc=sem_cc,
    )

    total_count = count_cont_invoices(**filter_kwargs)
    total_pages = max(1, math.ceil(total_count / PAGE_SIZE))
    page = min(page, total_pages)
    offset = (page - 1) * PAGE_SIZE

    invoices = get_cont_invoices(
        **filter_kwargs,
        order_by=order_by,
        order_dir=order_dir,
        limit=PAGE_SIZE,
        offset=offset,
    )

    # Summary cards (always global, not filtered)
    try:
        summary = get_cont_summary()
    except Exception:
        summary = {'por_contabilizar_count': 0, 'por_contabilizar_eur': 0,
                   'contabilizado_mes': 0, 'tickets_abertos': 0}

    stores = get_stores_list()
    all_supplier_names = get_distinct_supplier_names()

    has_filters = bool(
        supplier_name or document_type or accounting_status
        or date_from_raw or date_to_raw or search or centro_custo_filter
    )

    return render_template(
        'contabilidade/index.html',
        invoices=invoices,
        summary=summary,
        stores=stores,
        all_supplier_names=all_supplier_names,
        cost_centers=get_cost_centers(ativo_only=False),
        supplier_name=supplier_name,
        centro_custo_filter=centro_custo_filter,
        document_type=document_type,
        accounting_status=accounting_status,
        date_from_raw=date_from_raw,
        date_to_raw=date_to_raw,
        search=search,
        order_by=order_by,
        order_dir=order_dir,
        page=page,
        total_pages=total_pages,
        total_count=total_count,
        page_size=PAGE_SIZE,
        has_filters=has_filters,
        accounting_status_labels=ACCOUNTING_STATUS_LABELS,
        document_type_labels=DOCUMENT_TYPE_LABELS,
        today=today,
    )


# ── Update accounting status (AJAX) ──────────────────────────────────────────

@contabilidade_bp.route('/atualizar-estado', methods=['POST'])
@perm_required('acesso_contabilidade')
def atualizar_estado():
    from db.contabilidade import (
        update_accounting_status, bulk_update_accounting_status,
        ACCOUNTING_STATUS_LABELS,
    )
    data = request.get_json(silent=True) or {}
    status = (data.get('status') or '').strip()
    if status not in ACCOUNTING_STATUS_LABELS:
        return jsonify({'ok': False, 'error': 'Estado inválido'}), 400

    invoice_ids = data.get('invoice_ids', [])
    notes = data.get('notes', None)
    username = _username()

    if not invoice_ids:
        return jsonify({'ok': False, 'error': 'Nenhuma fatura selecionada'}), 400

    if len(invoice_ids) == 1:
        ok = update_accounting_status(invoice_ids[0], status, username, notes=notes)
        updated = 1 if ok else 0
    else:
        updated = bulk_update_accounting_status(invoice_ids, status, username)

    return jsonify({
        'ok': True,
        'updated': updated,
        'status': status,
        'status_label': ACCOUNTING_STATUS_LABELS[status],
    })


# ── Selected invoice ZIP download ─────────────────────────────────────────────

@contabilidade_bp.route('/faturas/zip', methods=['POST'])
@perm_required('acesso_contabilidade')
def download_selected_zip():
    from db.contabilidade import (
        get_cont_invoice_zip_data,
        get_cont_invoice_zip_metadata,
    )

    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({'ok': False, 'error': 'Pedido de ZIP inválido.'}), 400

    raw_ids = data.get('invoice_ids')
    if not isinstance(raw_ids, list) or not raw_ids:
        return jsonify({
            'ok': False,
            'error': 'Selecione pelo menos uma fatura.',
        }), 400
    if len(raw_ids) > MAX_CONTABILIDADE_ZIP_DOCUMENTS:
        return jsonify({
            'ok': False,
            'error': (
                f'O ZIP pode conter no máximo '
                f'{MAX_CONTABILIDADE_ZIP_DOCUMENTS} faturas.'
            ),
        }), 413

    invoice_ids = []
    seen_ids = set()
    for raw_id in raw_ids:
        if isinstance(raw_id, bool):
            return jsonify({
                'ok': False,
                'error': 'A seleção contém IDs de fatura inválidos.',
            }), 400
        if isinstance(raw_id, int):
            invoice_id = raw_id
        elif isinstance(raw_id, str) and raw_id.isascii() and raw_id.isdecimal():
            try:
                invoice_id = int(raw_id)
            except ValueError:
                return jsonify({
                    'ok': False,
                    'error': 'A seleção contém IDs de fatura inválidos.',
                }), 400
        else:
            return jsonify({
                'ok': False,
                'error': 'A seleção contém IDs de fatura inválidos.',
            }), 400
        if invoice_id <= 0 or invoice_id > (2**63 - 1) or invoice_id in seen_ids:
            return jsonify({
                'ok': False,
                'error': 'A seleção contém IDs inválidos ou repetidos.',
            }), 400
        invoice_ids.append(invoice_id)
        seen_ids.add(invoice_id)

    try:
        zip_folder_name = _safe_zip_folder_name(data.get('zip_name'))
    except ValueError as exc:
        return jsonify({'ok': False, 'error': str(exc)}), 400

    try:
        metadata_records = get_cont_invoice_zip_metadata(invoice_ids)
    except Exception:
        logger.exception('Falha ao validar documentos para ZIP da Contabilidade.')
        return jsonify({
            'ok': False,
            'error': 'Não foi possível verificar as faturas selecionadas.',
        }), 500

    metadata_by_id = {int(record['id']): record for record in metadata_records}
    if set(metadata_by_id) != set(invoice_ids):
        return jsonify({
            'ok': False,
            'error': (
                'Uma ou mais faturas são inválidas ou não pertencem '
                'à lista da Contabilidade.'
            ),
        }), 400

    total_metadata_bytes = 0
    for invoice_id in invoice_ids:
        file_size = metadata_by_id[invoice_id].get('pdf_size')
        if not file_size or file_size <= 0:
            return jsonify({
                'ok': False,
                'error': (
                    'Uma ou mais faturas selecionadas não têm documento; '
                    'não foi criado ZIP.'
                ),
            }), 422
        total_metadata_bytes += file_size
        if total_metadata_bytes > MAX_CONTABILIDADE_ZIP_BYTES:
            return jsonify({
                'ok': False,
                'error': 'O tamanho total dos documentos excede o limite do ZIP (50 MB).',
            }), 413

    try:
        data_records = get_cont_invoice_zip_data(invoice_ids)
    except Exception:
        logger.exception('Falha ao carregar documentos para ZIP da Contabilidade.')
        return jsonify({
            'ok': False,
            'error': 'Não foi possível carregar as faturas selecionadas.',
        }), 500

    data_by_id = {int(record['id']): record for record in data_records}
    if set(data_by_id) != set(invoice_ids):
        return jsonify({
            'ok': False,
            'error': (
                'Uma ou mais faturas deixaram de pertencer à lista da Contabilidade; '
                'não foi criado ZIP.'
            ),
        }), 400

    files = []
    total_source_bytes = 0
    for invoice_id in invoice_ids:
        record = metadata_by_id[invoice_id]
        raw_data = data_by_id[invoice_id].get('pdf_data')
        if not raw_data:
            return jsonify({
                'ok': False,
                'error': (
                    'Uma ou mais faturas selecionadas não têm documento; '
                    'não foi criado ZIP.'
                ),
            }), 422
        if isinstance(raw_data, memoryview):
            file_data = raw_data.tobytes()
        elif isinstance(raw_data, (bytes, bytearray)):
            file_data = bytes(raw_data)
        else:
            logger.error('Documento com tipo inválido ao preparar ZIP da Contabilidade.')
            return jsonify({
                'ok': False,
                'error': 'Não foi possível ler um dos documentos selecionados.',
            }), 500
        if not file_data:
            return jsonify({
                'ok': False,
                'error': (
                    'Uma ou mais faturas selecionadas não têm documento; '
                    'não foi criado ZIP.'
                ),
            }), 422
        total_source_bytes += len(file_data)
        if total_source_bytes > MAX_CONTABILIDADE_ZIP_BYTES:
            return jsonify({
                'ok': False,
                'error': 'O tamanho total dos documentos excede o limite do ZIP (50 MB).',
            }), 413
        files.append((invoice_id, record, file_data))

    files.sort(key=lambda item: item[0])
    archive_buffer = BytesIO()
    used_names = set()
    try:
        with zipfile.ZipFile(
            archive_buffer, mode='w', compression=zipfile.ZIP_DEFLATED, compresslevel=6
        ) as archive:
            for invoice_id, record, file_data in files:
                mimetype = _detect_invoice_mime(file_data, record.get('pdf_filename'))
                filename = _safe_invoice_download_name(
                    record.get('pdf_filename'),
                    mimetype,
                    record.get('cost_center_names') or (),
                )
                filename = _unique_zip_entry_filename(filename, invoice_id, used_names)
                archive.writestr(
                    f'{zip_folder_name}/{filename}',
                    file_data,
                )
    except Exception:
        logger.exception('Falha ao criar ZIP de faturas da Contabilidade.')
        return jsonify({
            'ok': False,
            'error': 'Não foi possível criar o ZIP das faturas selecionadas.',
        }), 500

    if archive_buffer.tell() > MAX_CONTABILIDADE_ZIP_BYTES:
        return jsonify({
            'ok': False,
            'error': 'O tamanho do ZIP excede o limite de 50 MB.',
        }), 413

    archive_buffer.seek(0)
    download_name = f'{zip_folder_name}.zip'
    response = send_file(
        archive_buffer,
        mimetype='application/zip',
        as_attachment=True,
        download_name=download_name,
        max_age=0,
    )
    response.headers['X-Download-Filename'] = quote(download_name, safe='')
    response.headers['Cache-Control'] = 'no-store'
    return response


# ── PDF download (reuse compras pattern) ─────────────────────────────────────

@contabilidade_bp.route('/fatura/<int:invoice_id>/pdf')
@perm_required('acesso_contabilidade')
def download_pdf(invoice_id: int):
    from db.faturas import get_invoice_pdf

    pdf_data, pdf_filename = get_invoice_pdf(invoice_id)
    if not pdf_data:
        return 'Ficheiro não disponível', 404

    as_attachment = request.args.get('dl') == '1'
    mimetype = _detect_invoice_mime(pdf_data, pdf_filename)
    cost_center_names = []
    if as_attachment:
        from db.contabilidade import get_cont_invoice_cost_center_names
        cost_center_names = get_cont_invoice_cost_center_names(invoice_id)
    filename = _safe_invoice_download_name(
        pdf_filename, mimetype, cost_center_names if as_attachment else ()
    )
    response = send_file(
        BytesIO(bytes(pdf_data) if isinstance(pdf_data, memoryview) else pdf_data),
        mimetype=mimetype,
        as_attachment=as_attachment,
        download_name=filename,
    )
    return response


@contabilidade_bp.route('/fatura/<int:invoice_id>/detalhe')
@perm_required('acesso_contabilidade')
def detalhe_fatura(invoice_id: int):
    """Render a read-only invoice detail and activity fragment for Contabilidade."""
    from db.faturas import (
        DOCUMENT_TYPE_LABELS,
        get_invoice,
        get_invoice_audit_log,
        get_invoice_status_labels_map,
    )
    from db.contabilidade import ACCOUNTING_STATUS_LABELS, ACCOUNTING_STATUS_BADGE

    inv = get_invoice(invoice_id)
    if not inv:
        abort(404)

    accounting_status = inv.get('accounting_status') or 'por_contabilizar'
    inv['accounting_status'] = accounting_status
    inv['accounting_status_label'] = ACCOUNTING_STATUS_LABELS.get(
        accounting_status, accounting_status
    )
    inv['accounting_status_badge'] = ACCOUNTING_STATUS_BADGE.get(
        accounting_status, 'bg-secondary'
    )

    return render_template(
        'contabilidade/_invoice_panel.html',
        inv=inv,
        audit_log=get_invoice_audit_log(invoice_id) or [],
        document_type_labels=DOCUMENT_TYPE_LABELS,
        status_labels=get_invoice_status_labels_map(),
        accounting_status_labels=ACCOUNTING_STATUS_LABELS,
    )


# ── Excel export ──────────────────────────────────────────────────────────────

@contabilidade_bp.route('/export.xlsx')
@perm_required('acesso_contabilidade')
def export_xlsx():
    from db.contabilidade import get_cont_invoices, ACCOUNTING_STATUS_LABELS
    from db.faturas import DOCUMENT_TYPE_LABELS

    supplier_name = request.args.get('supplier_name', '').strip()
    document_type = request.args.get('document_type', '').strip()
    if document_type not in DOCUMENT_TYPE_LABELS:
        document_type = ''
    accounting_status = request.args.get('accounting_status', '').strip()
    if accounting_status not in ACCOUNTING_STATUS_LABELS and accounting_status != '':
        accounting_status = ''
    date_from = _parse_date(request.args.get('date_from', ''))
    date_to = _parse_date(request.args.get('date_to', ''))
    search = request.args.get('q', '').strip()
    centro_custo_id, sem_cc, _centro_custo_filter = (
        _parse_centro_custo_filter(request.args.get('centro_custo_id', ''))
    )

    rows = get_cont_invoices(
        supplier_name=supplier_name or None,
        document_type=document_type or None,
        accounting_status=accounting_status or None,
        date_from=date_from,
        date_to=date_to,
        search=search or None,
        centro_custo_id=centro_custo_id,
        sem_cc=sem_cc,
        limit=10000,
        offset=0,
        order_by='issue_date',
        order_dir='desc',
    )

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return 'openpyxl não disponível', 500

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Contabilidade'

    headers = [
        'Data Emissão', 'Fornecedor', 'NIF', 'Centro de Custo',
        'Nº Documento', 'Tipo',
        'Valor (€)', 'IVA (€)', 'Estado Pagamento',
        'Estado Contabilístico', 'Notas Contabilidade',
        'Data Contabilização', 'Contabilizado por',
    ]
    header_fill = PatternFill('solid', fgColor='1F4E79')
    header_font = Font(color='FFFFFF', bold=True)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    STATUS_LABELS = {
        'pending_review': 'A rever',
        'scheduled': 'Agendado',
        'paid': 'Pago',
        'cancelled': 'Cancelado',
    }

    for r_idx, inv in enumerate(rows, 2):
        ws.cell(r_idx, 1, inv['issue_date'].strftime('%d/%m/%Y') if inv['issue_date'] else '')
        ws.cell(r_idx, 2, inv.get('supplier_display_name') or inv['supplier_name'] or '')
        ws.cell(r_idx, 3, inv['supplier_nif'] or '')
        ws.cell(r_idx, 4, inv.get('centro_custo_name') or 'Sem centro de custo')
        ws.cell(r_idx, 5, inv['invoice_number'] or '')
        ws.cell(r_idx, 6, DOCUMENT_TYPE_LABELS.get(inv['document_type'] or 'fatura', inv['document_type'] or ''))
        ws.cell(r_idx, 7, float(inv['amount_eur']) if inv['amount_eur'] is not None else '')
        ws.cell(r_idx, 8, float(inv['vat_amount_eur']) if inv['vat_amount_eur'] is not None else '')
        ws.cell(r_idx, 9, STATUS_LABELS.get(inv['status'], inv['status'] or ''))
        ws.cell(r_idx, 10, ACCOUNTING_STATUS_LABELS.get(inv['accounting_status'], inv['accounting_status'] or ''))
        ws.cell(r_idx, 11, inv['accounting_notes'] or '')
        ws.cell(r_idx, 12, inv['accounting_updated_at'].strftime('%d/%m/%Y %H:%M') if inv['accounting_updated_at'] else '')
        ws.cell(r_idx, 13, inv['accounting_updated_by'] or '')

    # Auto column width
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    today_str = _date.today().strftime('%Y%m%d')
    filename = f'contabilidade_{today_str}.xlsx'
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


# ── Tickets ───────────────────────────────────────────────────────────────────

@contabilidade_bp.route('/tickets')
@perm_required('acesso_contabilidade')
def tickets():
    from db.contabilidade import get_tickets, TICKET_STATUS_LABELS, get_cont_summary

    status_filter = request.args.get('status', '').strip()
    if status_filter not in TICKET_STATUS_LABELS and status_filter != '':
        status_filter = ''

    ticket_list = get_tickets(status=status_filter or None)

    try:
        summary = get_cont_summary()
    except Exception:
        summary = {'tickets_abertos': 0}

    return render_template(
        'contabilidade/tickets.html',
        tickets=ticket_list,
        status_filter=status_filter,
        ticket_status_labels=TICKET_STATUS_LABELS,
        summary=summary,
        today=_date.today(),
    )


@contabilidade_bp.route('/tickets/criar', methods=['POST'])
@perm_required('acesso_contabilidade')
def criar_ticket():
    from db.contabilidade import create_ticket

    titulo = request.form.get('titulo', '').strip()
    descricao = request.form.get('descricao', '').strip()
    prazo_raw = request.form.get('prazo', '').strip()
    invoice_id_raw = request.form.get('invoice_id', '').strip()
    invoice_required = request.form.get('invoice_id_required') == '1'
    return_url_raw = request.form.get('_return_url', '').strip()
    return_url = _safe_contabilidade_return_url(return_url_raw)
    list_return = url_for('contabilidade.index') if invoice_required else None
    error_return = return_url or list_return or url_for('contabilidade.tickets')

    if not titulo:
        flash('O título do ticket é obrigatório.', 'error')
        return redirect(error_return)

    prazo = _parse_date(prazo_raw)
    invoice_id = None
    if invoice_id_raw:
        if not invoice_id_raw.isdigit() or int(invoice_id_raw) <= 0:
            flash('A fatura associada é inválida. O ticket não foi criado.', 'error')
            return redirect(error_return)
        invoice_id = int(invoice_id_raw)
        from db.faturas import get_invoice
        if not get_invoice(invoice_id):
            flash('A fatura selecionada já não existe. O ticket não foi criado.', 'error')
            return redirect(error_return)
    elif invoice_required:
        flash('Selecione uma fatura antes de criar o ticket.', 'error')
        return redirect(error_return)

    ticket_id = create_ticket(
        titulo=titulo,
        descricao=descricao,
        criado_por=_username(),
        invoice_id=invoice_id,
        prazo=prazo,
    )
    flash('Ticket criado com sucesso.', 'success')

    if return_url or invoice_required:
        return redirect(return_url or url_for('contabilidade.index'))

    return redirect(url_for('contabilidade.ticket_detalhe', ticket_id=ticket_id))


@contabilidade_bp.route('/tickets/<int:ticket_id>')
@perm_required('acesso_contabilidade')
def ticket_detalhe(ticket_id: int):
    from db.contabilidade import get_ticket, TICKET_STATUS_LABELS

    ticket = get_ticket(ticket_id)
    if not ticket:
        flash('Ticket não encontrado.', 'error')
        return redirect(url_for('contabilidade.tickets'))

    return render_template(
        'contabilidade/ticket_detalhe.html',
        ticket=ticket,
        ticket_status_labels=TICKET_STATUS_LABELS,
    )


@contabilidade_bp.route('/tickets/<int:ticket_id>/anexar-ficheiro', methods=['POST'])
@perm_required('acesso_contabilidade')
def anexar_ficheiro_ticket(ticket_id: int):
    from db.contabilidade import get_ticket, add_ticket_response
    from db.faturas import save_invoice_pdf

    ticket = get_ticket(ticket_id)
    if not ticket:
        flash('Ticket não encontrado.', 'error')
        return redirect(url_for('contabilidade.tickets'))

    if not ticket.get('invoice_id'):
        flash('Este ticket não está associado a uma fatura.', 'error')
        return redirect(url_for('contabilidade.ticket_detalhe', ticket_id=ticket_id))

    pdf_file = request.files.get('pdf_file')
    if not pdf_file or not pdf_file.filename:
        flash('Nenhum ficheiro seleccionado.', 'warning')
        return redirect(url_for('contabilidade.ticket_detalhe', ticket_id=ticket_id))

    ext = pdf_file.filename.rsplit('.', 1)[-1].lower() if '.' in pdf_file.filename else ''
    if ext != 'pdf':
        flash('Apenas ficheiros PDF são aceites.', 'warning')
        return redirect(url_for('contabilidade.ticket_detalhe', ticket_id=ticket_id))

    pdf_data = pdf_file.read()
    if not pdf_data.startswith(b'%PDF'):
        flash('Ficheiro não é um PDF válido.', 'warning')
        return redirect(url_for('contabilidade.ticket_detalhe', ticket_id=ticket_id))

    save_invoice_pdf(ticket['invoice_id'], pdf_data, pdf_file.filename or 'fatura.pdf')

    add_ticket_response(
        ticket_id=ticket_id,
        mensagem='Ficheiro anexado',
        novo_status='resolvido',
        criado_por=_username(),
    )

    flash('Ficheiro anexado e ticket resolvido com sucesso.', 'success')
    return redirect(url_for('contabilidade.ticket_detalhe', ticket_id=ticket_id))


@contabilidade_bp.route('/tickets/<int:ticket_id>/responder', methods=['POST'])
@perm_required('acesso_contabilidade')
def responder_ticket(ticket_id: int):
    from db.contabilidade import add_ticket_response, TICKET_STATUS_LABELS

    mensagem = request.form.get('mensagem', '').strip()
    novo_status = request.form.get('novo_status', '').strip()

    if not mensagem:
        flash('A mensagem de resposta não pode estar vazia.', 'error')
        return redirect(url_for('contabilidade.ticket_detalhe', ticket_id=ticket_id))

    if novo_status and novo_status not in TICKET_STATUS_LABELS:
        novo_status = ''

    add_ticket_response(
        ticket_id=ticket_id,
        mensagem=mensagem,
        novo_status=novo_status or None,
        criado_por=_username(),
    )
    flash('Resposta adicionada.', 'success')
    return redirect(url_for('contabilidade.ticket_detalhe', ticket_id=ticket_id))
